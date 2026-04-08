import os
from typing import List
from pypdf import PdfReader
from docx import Document as DocxDocument
from pptx import Presentation
from openpyxl import load_workbook

class DocumentProcessor:
    def __init__(self, upload_dir: str = "uploads"):
        self.upload_dir = upload_dir
        os.makedirs(upload_dir, exist_ok=True)

    def _extract_pdf_with_pages(self, file_path: str) -> List[dict]:
        results = []
        try:
            reader = PdfReader(file_path)
            for page_num, page in enumerate(reader.pages, start=1):
                text = page.extract_text()
                if text and text.strip():
                    results.append({"page": page_num, "text": text})
        except Exception as e:
            print(f"Error extracting PDF with pages: {e}")
        return results

    def _extract_pptx_with_pages(self, file_path: str) -> List[dict]:
        results = []
        try:
            prs = Presentation(file_path)
            for slide_num, slide in enumerate(prs.slides, start=1):
                text_parts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text_parts.append(shape.text)
                text = "\n".join(text_parts)
                if text.strip():
                    results.append({"page": slide_num, "text": text})
        except Exception as e:
            print(f"Error extracting PPTX with pages: {e}")
        return results

    def _extract_docx(self, file_path: str) -> str:
        try:
            doc = DocxDocument(file_path)
            return "\n".join([para.text for para in doc.paragraphs])
        except Exception as e:
            print(f"Error extracting DOCX: {e}")
            return ""

    def _extract_xlsx(self, file_path: str) -> str:
        try:
            wb = load_workbook(file_path, data_only=True)
            text = ""
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                text += f"\nSheet: {sheet_name}\n"
                for row in sheet.iter_rows(values_only=True):
                    row_text = " | ".join([str(cell) if cell else "" for cell in row])
                    if row_text.strip():
                        text += row_text + "\n"
            return text
        except Exception as e:
            print(f"Error extracting XLSX: {e}")
            return ""

    def extract_text(self, file_path: str, file_type: str) -> List[dict]:
        if file_type == "pdf":
            return self._extract_pdf_with_pages(file_path)
        elif file_type == "docx":
            text = self._extract_docx(file_path)
            return [{"page": 1, "text": text}] if text.strip() else []
        elif file_type == "pptx":
            return self._extract_pptx_with_pages(file_path)
        elif file_type == "xlsx":
            text = self._extract_xlsx(file_path)
            return [{"page": 1, "text": text}] if text.strip() else []
        return []

    def chunk_text(self, text: str, chunk_size: int = 500, overlap: int = 50) -> List[str]:
        words = text.split()
        chunks = []
        
        for i in range(0, len(words), chunk_size - overlap):
            chunk = " ".join(words[i:i + chunk_size])
            if chunk.strip():
                chunks.append(chunk)
        
        return chunks

    def chunk_text_with_pages(self, pages: List[dict], chunk_size: int = 500) -> List[dict]:
        chunks = []
        for page_info in pages:
            page = page_info["page"]
            text = page_info["text"]
            words = text.split()
            for i in range(0, len(words), chunk_size):
                chunk_words = words[i:i + chunk_size]
                chunks.append({
                    "page": page,
                    "chunk_index": len(chunks),
                    "text": " ".join(chunk_words)
                })
        return chunks

document_processor = DocumentProcessor()
